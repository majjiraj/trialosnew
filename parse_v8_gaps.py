import openpyxl, json
from pathlib import Path
p=Path('/Users/rajeshmajji/trialo/feedback/USDM_Gap_Analysis_B7981027_v8.xlsx')
wb=openpyxl.load_workbook(p,data_only=True)
print('sheets',wb.sheetnames)
ws=wb[wb.sheetnames[0]]
rows=list(ws.iter_rows(values_only=True))
header_row=None
for i,r in enumerate(rows,1):
    if r and str(r[0]).strip()=='#':
        header_row=i; break
print('header_row',header_row)
if header_row:
    h=[str(c).strip() if c is not None else '' for c in rows[header_row-1]]
    print('headers',h)
    items=[]
    for r in rows[header_row:]:
        if r and str(r[0]).strip().isdigit():
            d={h[j]:r[j] for j in range(min(len(h),len(r)))}
            items.append(d)
    print('TOTAL_ITEMS',len(items))
    for d in items:
        print(json.dumps({
            'num':d.get('#'),
            'section':d.get('USDM Section'),
            'field':d.get('Class / Field'),
            'gap':d.get('Gap Description'),
            'severity':d.get('Severity'),
            'status':d.get('Status'),
            'rec':d.get('Recommendation / Suggested Action')
        }, default=str))
wb.close()
